"""Read-only result snapshots; CSV display escaping never changes catalog values.

The deployed Python-only environment has no Codex artifact-tool runtime.
XLSX uses openpyxl as the portable export fallback, not for dataset authoring.
"""
import csv
from io import BytesIO,StringIO
import json
from catalog.repository import CatalogRepository
from catalog.evidence import verify_bundle
from catalog.binding import verify_result


def _safe_cell(value):
    if isinstance(value,str) and (value.lstrip(' \t\r\n').startswith(('=','+','-','@')) or value.startswith(('\t','\r','\n'))):return "'"+value
    return value


def result_records(result,bundle):
    verify_bundle(bundle)
    verify_result(result)
    if result.decision_sha256!=bundle['decision_sha256'] or result.ranking_sha256!=bundle['ranking_sha256']:
        raise ValueError('导出排名与证据不是同一次决策')
    for key in ('dataset_version','feature_rule_version','algorithm_version'):
        if getattr(result,key)!=bundle[key]:raise ValueError('结果与证据版本不一致')
    rows=[]
    for original in CatalogRepository.as_records(result.ranked):
        row={k:original.get(k) for k in ('sku_id','model_id','variant_id','display_name','variant_name','ram_raw','storage_raw','price_id','price_cny','condition','observed_at','is_estimated','score')}
        for fact in bundle['facts'].values():
            if fact['sku_id']==row['sku_id']:row[fact['metric_key']]=fact['value']
        row.update({k:bundle[k] for k in ('dataset_version','feature_rule_version','algorithm_version')})
        context=bundle['decision_context']
        row.update(target_screen=context['target_screen'],decision_sha256=bundle['decision_sha256'],
            ranking_sha256=bundle['ranking_sha256'],
            weights_json=json.dumps(context['weights'],ensure_ascii=False,allow_nan=False),
            selection_rules_json=json.dumps(context['selection_rules'],ensure_ascii=False,allow_nan=False))
        rows.append(row)
    return rows


def export_results(result,bundle,format: str) -> bytes:
    rows=result_records(result,bundle)
    if format=='json':
        return json.dumps({'ranking':rows,'backups':CatalogRepository.as_records(result.backups),'evidence':bundle},ensure_ascii=False,allow_nan=False,indent=2).encode('utf-8')
    columns=list(dict.fromkeys(k for row in rows for k in row)) or ['sku_id']
    if format=='csv':
        stream=StringIO(newline='');writer=csv.DictWriter(stream,fieldnames=columns);writer.writeheader()
        writer.writerows({k:_safe_cell(v) for k,v in row.items()} for row in rows)
        return stream.getvalue().encode('utf-8-sig')
    if format!='xlsx':raise ValueError('只支持 csv、json、xlsx')
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='推荐结果';ws.append(columns)
    for row in rows:ws.append([_safe_cell(row.get(k)) for k in columns])
    meta=wb.create_sheet('证据与说明');meta.append(['项目','内容'])
    for key in ('dataset_version','feature_rule_version','algorithm_version','evidence_sha256'):meta.append([key,bundle[key]])
    for key,value in bundle['decision_context'].items():
        text=json.dumps(value,ensure_ascii=False,allow_nan=False) if isinstance(value,(dict,list)) else value
        if isinstance(text,str) and len(text)>160:
            for start in range(0,len(text),160):meta.append([key+f' [{start//160+1}]',text[start:start+160]])
        else:meta.append([key,text])
    meta.append(['展示导出规则','缺失为留空；公式型文本前加单引号，只改变展示副本，不修改数据库。'])
    for text in bundle['limitations']:meta.append(['限制',text])
    facts=wb.create_sheet('逐项事实');facts.append(['事实 ID','字段','内容'])
    for fact in bundle['facts'].values():
        for key,value in fact.items():
            text=json.dumps(value,ensure_ascii=False,allow_nan=False) if isinstance(value,(dict,list)) else value
            if isinstance(text,str) and len(text)>160:
                for start in range(0,len(text),160):
                    facts.append([fact['fact_id'],f'{key} [{start//160+1}]',_safe_cell(text[start:start+160])])
            else:facts.append([fact['fact_id'],key,_safe_cell(text)])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(col)].width=24 if sheet==ws else (34 if col==1 else (28 if sheet==facts and col==2 else 85))
        for cell in sheet[1]:cell.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173754')
        for line in list(sheet.rows)[1:]:
            sheet.row_dimensions[line[0].row].height=34 if sheet==ws else 66
            for cell in line:
                cell.font=Font(name='Microsoft YaHei',size=10);cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,str):cell.data_type='s';cell.number_format='@'
                elif isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):cell.number_format='0.####'
    stream=BytesIO();wb.save(stream);wb.close();return stream.getvalue()
