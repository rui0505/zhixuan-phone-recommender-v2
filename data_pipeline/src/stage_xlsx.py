# -*- coding: utf-8 -*-
"""阶段六：规范 XLSX 导出（WPS 兼容）。

约定：
- 标准 .xlsx，不含宏；ID 列写为文本并设 number_format='@'，避免科学计数法。
- 原文列标「只读原文」，派生列标单位与规则。
- 首行冻结 + 自动筛选，便于 WPS 内核对。
- 与 SQLite 同源导出，禁止人工维护两套值。
"""
from __future__ import annotations

import json
import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rules import FIELD_DICTIONARY, METRIC_DEFINITIONS, NON_COMPARABLE_METRICS, RULE_VERSION

TEXT_COLS = {"model_id", "variant_id", "sku_id", "price_id", "spec_id", "run_id",
             "observation_id", "asset_id", "color_id", "source_id", "source_record_id",
             "member_path", "cover_asset_id", "target_id", "sha256", "index_declared_sha256",
             "crc32", "asset_key", "issue_code", "metric_key", "spec_key", "registry_key"}

HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
HDR_FONT = Font(bold=True)


def _sheet(wb, title, header, rows, widths=None, note=None):
    if isinstance(wb, list):
        wb.append(dict(title=title, header=header, rows=rows, widths=widths or {}, note=note))
        return
    ws = wb.create_sheet(title)
    if note:
        ws.append([note])
        ws.cell(1, 1).font = Font(italic=True, color="7F7F7F")
        start = 2
    else:
        start = 1
    ws.append(header)
    for i, h in enumerate(header, start=1):
        c = ws.cell(start, i)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, start=1):
        letter = get_column_letter(i)
        w = (widths or {}).get(h)
        if w is None:
            sample = [len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None]
            w = min(max(sample or [8]) + 2, 46)
        ws.column_dimensions[letter].width = w
    if rows:
        for row in ws.iter_rows(min_row=start + 1, max_row=start + len(rows)):
            for cell in row:
                if cell.column_letter in {
                        get_column_letter(header.index(h) + 1) for h in TEXT_COLS if h in header}:
                    cell.number_format = "@"
    ws.freeze_panes = ws.cell(start + 1, 1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(header))}{start + len(rows)}"
    return ws


def export_workbook(path, cfg, manifest, entities, assets_pack, issues, counts, stats):
    # Codex 交付使用已验证的 artifact-tool；普通 Python 复现保留原导出器。
    # 两条路径共用下方同一份数据，不维护两份工作簿值。
    artifact_node = os.environ.get('DATASET_ARTIFACT_NODE')
    if artifact_node:
        wb = []
    else:
        wb = Workbook()
        wb.remove(wb.active)

    # ---- 数据说明 ----
    desc = [
        ["数据集版本", cfg["dataset_version"]],
        ["规则版本", RULE_VERSION],
        ["ID 注册表版本", cfg["id_registry_version"]],
        ["构建时间", manifest["generated_at"]],
        ["源目录", manifest["source_dir"]],
        ["输出目录", manifest["output_dir"]],
        ["原始层", "raw_rows / raw_cells，逐单元格保留值、值类型、公式、缓存值与数字格式"],
        ["原始文件只读", "构建前后 SHA-256 一致；未重算、未改写、未覆盖任何源文件"],
        ["派生约定", "1TB = 1024GB（项目换算约定，原始标称不变）；缺失一律 NULL，原因见「评测指标」"],
        ["价格口径", "condition 分 new / used / unknown；无成色证据不静默归类"],
        ["可比性", "comparison_group 表示来源+协议+条件组合；跨协议不直接混排"],
        ["本表用途", "WPS 查看、核验与团队补数；数据库为程序读取对象，两者同源导出"],
        ["未包含", "本版本不实现拍照硬件综合分，不拟合跨来源校准模型，不替换正式网页数据"],
    ]
    _sheet(wb, "数据说明", ["项目", "说明"], desc, widths={"项目": 20, "说明": 96})

    # ---- 机型 ----
    _sheet(wb, "机型",
           ["model_id", "品牌", "标准展示名称", "匹配键", "发布年月", "发布原文",
            "发布粒度", "首次出现源记录", "封面图片", "备注"],
           [[m["model_id"], m["brand"], m["display_name"], m["display_name_match_key"],
             m["release_month"], m["release_month_raw"], m["release_granularity"],
             m["first_source_record_id"], m["cover_asset_id"], m["note"]]
            for m in entities["phone_models"]],
           note="展示名称与发布原文来自源表，逐值保留；匹配键仅用于匹配，不改变原文。")

    # ---- 版本 ----
    _sheet(wb, "版本",
           ["variant_id", "model_id", "版本名称", "名称是否源文本", "地区", "地区依据",
            "卫星通信", "是否默认版本", "版本备注原文", "源记录"],
           [[v["variant_id"], v["model_id"], v["variant_name"], v["name_is_source"],
             v["region"], v["region_evidence"], v["satellite_support"], v["is_default"],
             v["variant_note_raw"], v["source_record_id"]]
            for v in entities["phone_variants"]],
           note="卫星通信版与标准版保持独立 ID，禁止自动合并；未知地区保留 NULL。")

    # ---- 配置 ----
    _sheet(wb, "配置",
           ["sku_id", "variant_id", "model_id", "运行内存(只读原文)", "ram_gb", "存储(只读原文)",
            "storage_gb", "内存类型", "内存类型依据", "源记录"],
           [[s["sku_id"], s["variant_id"], s["model_id"], s["ram_raw"], s["ram_gb"],
             s["storage_raw"], s["storage_gb"], s["memory_type"], s["memory_type_source"],
             s["source_record_id"]] for s in entities["skus"]],
           note="ram_gb/storage_gb 为派生值；原文与单位同时保留。1TB(LPDDR5T) 与 1TB 为不同配置。")

    # ---- 价格快照 ----
    _sheet(wb, "价格快照",
           ["price_id", "sku_id", "价格(只读原文)", "price_cny", "币种", "价格类型",
            "成色", "是否估计值", "渠道", "卖家/来源原文", "价格说明原文", "限定条件",
            "原始销售状态", "采集日期", "日期口径", "来源ID", "源记录"],
           [[p["price_id"], p["sku_id"], p["price_raw"], p["amount"], p["currency"],
             p["price_kind"], p["condition"], p["is_estimated"], p["channel"], p["seller"],
             p["price_note"], p["qualifiers"], p["source_sale_status"], p["observed_at"],
             p["observed_at_kind"], p["source_id"], p["source_record_id"]]
            for p in entities["price_snapshots"]],
           note="原始停售状态不等于渠道无库存；二手约价标记 estimated；区间价不取首值。")

    # ---- 规格参数 ----
    _sheet(wb, "规格参数",
           ["spec_id", "适用对象类型", "适用对象ID", "variant_id", "字段", "单位",
            "派生值", "文本值", "原始值(只读原文)", "原始值类型", "源列", "解析状态",
            "口径/条件说明", "可比性分组", "来源ID", "源记录", "规则版本"],
           [[s["spec_id"], s["subject_type"], s["subject_id"], s["variant_id"], s["spec_key"],
             s["unit"], s["value_num"], s["value_text"], s["raw_value"], s["raw_value_type"],
             s["source_col"], s["parse_status"], s["note"], s["comparison_group"], s["source_id"],
             s["source_record_id"], s["rule_version"]]
            for s in entities["specifications"]],
           note="SKU 优先；仅当前全部配置一致的规格提升为机型共有值。无法精确解析时保持 NULL。"
                "「口径/条件说明」保留源表原文；「可比性分组」为规范化分组键，仅用于比较，"
                "不同组不直接混排。")

    _sheet(wb, '规格来源映射', ['spec_id','source_record_id','源列','source_id'],
           [[p['spec_id'],p['source_record_id'],p['source_col'],p['source_id']] for p in entities.get('spec_provenance',[])])

    # ---- 评测运行 ----
    _sheet(wb, "评测运行",
           ["run_id", "对象类型", "对象ID", "model_id", "variant_id", "sku_id", "项目",
            "机构", "协议", "协议/条件原文", "结构化条件", "日期", "系统版本",
            "可比性分组", "适用范围", "范围说明", "来源ID", "源记录", "数据备注"],
           [[r["run_id"], r["subject_type"], r["subject_id"], r["model_id"], r["variant_id"],
             r["sku_id"], r["benchmark"], r["institution"], r["protocol"],
             r["protocol_detail"], r["conditions_json"], r["run_date"], r["os_version"],
             r["comparison_group"], r["applicable_scope"], r["scope_note"], r["source_id"],
             r["source_record_id"], r["data_note"]] for r in entities["benchmark_runs"]],
           note="一次测试一条记录；机型级共享结果已标注适用范围，不逐配置重复。")

    # ---- 评测指标 ----
    _sheet(wb, "评测指标",
           ["observation_id", "run_id", "指标", "数值", "单位", "原始值(只读原文)",
            "原始值类型", "缺失原因", "可比性分组", "源列", "核验状态", "备注", "源记录"],
           [[o["observation_id"], o["run_id"], o["metric_key"], o["value_num"], o["unit"],
             o["raw_value"], o["raw_value_type"], o["missing_reason"], o["comparison_group"],
             o["source_col"], o["verification_status"], o["note"], o["source_record_id"]]
            for o in entities["metric_observations"]],
           note="缺失原因区分 not_found（未查到）与 not_reported（未公开）；不填 0、不推算实测。")

    # ---- 来源 ----
    _sheet(wb, "来源",
           ["source_id", "来源类别", "原始来源文字(只读原文)", "可解析URL", "URL数量",
            "发布方", "资料日期", "日期粒度", "采集日期", "有可解析链接"],
           [[s["source_id"], s["source_kind"], s["raw_text"], s["url"], s["url_count"],
             s["publisher"], s["published_at"], s["published_granularity"],
             s["collected_at"], s["url_is_parseable"]]
            for s in entities["sources"]],
           note="无可解析 URL 的记录保留原始来源文字，不伪造链接。")

    # ---- 图片 ----
    _sheet(wb, "图片",
           ["asset_id", "model_id", "variant_id", "color_id", "颜色", "适用范围",
            "ZIP内原路径", "解压路径", "网页副本键", "格式", "宽", "高", "字节数", "SHA256",
            "解码状态", "索引声明SHA256", "哈希一致", "索引宽", "索引高", "尺寸一致",
            "来源类型", "来源页面", "原始地址", "是否封面", "备注"],
           [[a["asset_id"], a["model_id"], a["variant_id"], a["color_id"], a["color_name"],
             a["applicable_scope"], a["original_zip_path"], a["extracted_path"],
             a["derived_path"], a["format"], a["width"], a["height"], a["size_bytes"],
             a["sha256"], a["decode_status"], a["index_declared_sha256"], a["sha256_match"],
             a["index_width"], a["index_height"], a["dimension_match"], a["source_kind"],
             a["source_page"], a["source_url"], a["is_cover"], a["note"]]
            for a in assets_pack["assets"]],
           note="原图不覆盖不裁剪；网页副本另存 assets_derived/；解码失败不伪造替代图。")

    # ---- 评测摘录 ----
    _sheet(wb, "评测摘录",
           ["snippet_id", "source_id", "适用对象类型", "适用对象ID", "关联运行", "支持的指标",
            "摘录类型", "摘录原文(只读原文)", "源列", "来源列", "内容来源", "文字提取方式", "标签规则",
            "检索标签", "标签数", "备注", "源记录"],
           [[s["snippet_id"], s["source_id"], s["subject_type"], s["subject_id"],
             s["related_run_id"], s["related_metric_keys"], s["snippet_kind"],
             s["snippet_text"], s["source_col"], s["source_ref_col"], s["origin"],
             s["extraction_method"],
             s["tag_method"], s["tags"], s["tag_count"], s["note"], s["source_record_id"]]
            for s in entities.get("evidence_snippets", [])],
           note="摘录原文逐字保留；检索标签为 keyword_rule_v1 自动提取结果，"
                "与源表人工结论分开标记（见「内容来源」「文字提取方式」）。"
                "本任务不把描述性结论转成 0-10 评分。")

    # ---- 字段字典 ----
    dict_rows = [[f, layer, unit, desc] for f, layer, unit, desc in FIELD_DICTIONARY]
    for key, name_cn, unit, direction, prefix, vtype, rng in METRIC_DEFINITIONS:
        dict_rows.append([key, "指标定义", unit,
                          f"{name_cn}；方向={direction}；可比性前缀={prefix}；值类型={vtype}；"
                          f"范围检查={rng}；跨来源可比={0 if key in NON_COMPARABLE_METRICS else 1}"])
    _sheet(wb, "字段字典", ["字段/指标", "层级", "单位", "说明"], dict_rows)

    # ---- 关联映射 ----
    _sheet(wb, "关联映射",
           ["source_file_key", "工作表", "源行号", "source_record_id", "目标类型", "目标ID",
            "映射角色", "合并/待核验说明"],
           [[m["source_file_key"], m["sheet_name"], m["source_row_no"], m["source_record_id"],
             m["target_type"], m["target_id"], m["mapping_role"], m["merge_note"]]
            for m in entities["import_mapping"]],
           note="每条源记录的去向均可在此追溯。")

    # ---- 质检问题 ----
    _sheet(wb, "质检问题",
           ["issue_code", "严重度", "源记录", "说明"],
           [[i["issue_code"], i["severity"], i.get("source_record_id", ""), i["detail"]]
            for i in issues],
           note="无法确定或相互冲突的内容在此登记，处理继续，不猜值、不改原文件。")

    if artifact_node:
        payload = {'sheets': wb, 'text_cols': sorted(TEXT_COLS), 'output': path}
        payload_path = os.path.join(os.path.dirname(path), 'logs', 'workbook_payload.json')
        with open(payload_path, 'w', encoding='utf-8') as handle:
            json.dump(payload, handle, ensure_ascii=False)
        subprocess.run([artifact_node, os.path.join(os.path.dirname(__file__), 'export_workbook.mjs'),
                        payload_path], check=True)
    else:
        wb.save(path)
    return counts
