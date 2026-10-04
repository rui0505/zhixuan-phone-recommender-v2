# -*- coding: utf-8 -*-
"""阶段一：输入登记与原始快照层。

产出：
  source_manifest.json          输入哈希/工作表/行数/文件时间
  raw_snapshot/raw_rows.jsonl   行级原始快照
  raw_snapshot/raw_cells.jsonl  单元格级原始快照（值/类型/公式/缓存值/数字格式）
不修改任何源文件。
"""
from __future__ import annotations

import os

import openpyxl

from common import (file_stat, has_suspicious_encoding, now_iso, sha256_file,
                    value_to_text, value_type_of)

ROLES = {
    "price_main": "主输入（机型/配置/价格/规格/游戏补充）",
    "price_aux": "可选辅助核对文件（较早补充版，仅交叉比对，不合并导入）",
    "camera": "相机评测输入（DXOMARK/其他机构）",
}


def read_workbook(path: str, file_key: str):
    """返回 (sheet_meta, rows, cells)。公式层与缓存值层分别读取。"""
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    sheet_meta, rows, cells = [], [], []
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        header = [c.value for c in ws_f[1]]
        n_formula = 0
        for r in range(1, ws_f.max_row + 1):
            row_values, row_types = [], []
            for c in range(1, ws_f.max_column + 1):
                cf = ws_f.cell(r, c)
                cv = ws_v.cell(r, c)
                raw = cf.value
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                if formula:
                    n_formula += 1
                row_values.append(value_to_text(raw))
                row_types.append(value_type_of(raw))
                cells.append({
                    "source_record_id": f"{file_key}#{ws_f.title}#R{r}",
                    "file_key": file_key,
                    "sheet_name": ws_f.title,
                    "source_row_no": r,
                    "source_col_no": c,
                    "col_letter": cf.column_letter,
                    "field_name": header[c - 1] if c - 1 < len(header) else None,
                    "raw_value_text": value_to_text(raw),
                    "raw_value_type": value_type_of(raw),
                    "formula": formula or "",
                    "cached_value_text": value_to_text(cv.value) if formula else "",
                    "cached_value_type": value_type_of(cv.value) if formula else "empty",
                    "number_format": cf.number_format,
                })
            rows.append({
                "source_record_id": f"{file_key}#{ws_f.title}#R{r}",
                "file_key": file_key,
                "sheet_name": ws_f.title,
                "source_row_no": r,
                "is_header_row": r == 1,
                "values_text": row_values,
                "value_types": row_types,
            })
        sheet_meta.append({
            "name": ws_f.title,
            "max_row": ws_f.max_row,
            "max_col": ws_f.max_column,
            "data_rows": max(ws_f.max_row - 1, 0),
            "header": [value_to_text(h) if h is not None else None for h in header],
            "formula_cells": n_formula,
            "merged_ranges": [str(m) for m in ws_f.merged_cells.ranges],
            "freeze_panes": str(ws_f.freeze_panes) if ws_f.freeze_panes else None,
        })
    wb_f.close()
    wb_v.close()
    return sheet_meta, rows, cells


def build_manifest(cfg: dict, paths: dict):
    inputs = cfg["inputs"]
    files, all_rows, all_cells = [], [], []

    for key in ("price_main", "price_aux", "camera"):
        name = inputs[key]
        path = os.path.join(paths["source_dir"], name)
        if not os.path.exists(path):
            files.append({"file_key": key, "file_name": name, "abs_path": path,
                          "role": ROLES[key], "exists": False})
            continue
        st = file_stat(path)
        sheet_meta, rows, cells = read_workbook(path, key)
        suspicious = [{"sheet": c["sheet_name"], "row": c["source_row_no"],
                       "col": c["source_col_no"], "text": c["raw_value_text"][:80]}
                      for c in cells if has_suspicious_encoding(c["raw_value_text"])]
        files.append({
            "file_key": key,
            "file_name": name,
            "abs_path": path,
            "role": ROLES[key],
            "exists": True,
            "sha256_before": sha256_file(path),
            "size_bytes": st["size_bytes"],
            "mtime": st["mtime"],
            "sheets": sheet_meta,
            "row_count": sum(s["data_rows"] for s in sheet_meta),
            "col_count": sheet_meta[0]["max_col"] if sheet_meta else None,
            "encoding_scan": {"suspicious": bool(suspicious), "samples": suspicious[:20]},
        })
        all_rows.extend(rows)
        all_cells.extend(cells)

    zip_name = inputs["images_zip"]
    zip_path = os.path.join(paths["source_dir"], zip_name)
    if os.path.exists(zip_path):
        st = file_stat(zip_path)
        files.append({
            "file_key": "images_zip",
            "file_name": zip_name,
            "abs_path": zip_path,
            "role": "原始图片包",
            "exists": True,
            "sha256_before": sha256_file(zip_path),
            "size_bytes": st["size_bytes"],
            "mtime": st["mtime"],
            "sheets": [],
            "row_count": None,
            "col_count": None,
            "encoding_scan": {"suspicious": False, "samples": []},
        })

    manifest = {
        "generated_at": now_iso(),
        "rule_version": cfg["rule_version"],
        "source_dir": paths["source_dir"],
        "output_dir": paths["output_dir"],
        "files": files,
    }
    return manifest, all_rows, all_cells
