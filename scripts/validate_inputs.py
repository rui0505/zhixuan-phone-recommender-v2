from __future__ import annotations
import argparse
import hashlib
import json
import os
from pathlib import Path
import openpyxl


def validate_inputs(config_path: Path) -> dict:
    cfg = json.loads(Path(config_path).read_text(encoding='utf-8'))
    source = Path(os.environ.get('PHONE_SOURCE_DIR') or cfg['source_dir'])
    checks = []
    for key, name in cfg['inputs'].items():
        path = source / name
        item = {'file_key': key, 'status': 'PASS', 'sha256': None, 'reason': ''}
        if Path(name).name != name or not path.is_file():
            item.update(status='FAIL', reason='Input missing or not a basename')
        else:
            item['sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
            if path.suffix.lower() == '.xlsx':
                wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
                try:
                    sheet = cfg.get('camera_sheet') if key == 'camera' else cfg.get('primary_sheet')
                    if sheet not in wb.sheetnames:
                        item.update(status='FAIL', reason='Required sheet missing')
                    else:
                        headers = list(next(wb[sheet].iter_rows(max_row=1, values_only=True)))
                        missing = set(cfg.get('required_fields', {}).get(key, [])) - set(headers)
                        if missing:
                            item.update(status='FAIL', reason='Required fields missing: ' + ','.join(sorted(missing)))
                finally:
                    wb.close()
        checks.append(item)
    return {'status': 'PASS' if checks and all(c['status'] == 'PASS' for c in checks) else 'FAIL', 'checks': checks}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--config', type=Path, required=True)
    args = parser.parse_args()
    report = validate_inputs(args.config)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(0 if report['status'] == 'PASS' else 1)
