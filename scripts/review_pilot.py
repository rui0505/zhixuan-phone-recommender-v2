"""Machine comparison and evidence pack; never mark human/WPS checks as done."""
import argparse
import json
from pathlib import Path
import sqlite3
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'data_pipeline/src'))
from common import load_config, value_to_text


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--config', required=True)
    args = parser.parse_args()
    cfg = load_config(args.config)
    wb = openpyxl.load_workbook(Path(cfg['source_dir']) / cfg['inputs']['price_main'], data_only=False)
    ws = wb[cfg['primary_sheet']]
    headers = [c.value for c in ws[1]]
    db = ROOT / 'local_build/pilot/catalog.sqlite'
    con = sqlite3.connect(db.resolve().as_uri() + '?mode=ro', uri=True)
    con.row_factory = sqlite3.Row
    results = []
    for sku in con.execute('SELECT s.*,m.display_name FROM skus s JOIN phone_models m USING(model_id) ORDER BY s.sku_id'):
        source_row = int(sku['source_record_id'].rsplit('#R', 1)[1])
        raw = {key: ws.cell(source_row, i + 1).value for i, key in enumerate(headers)}
        identity = (sku['display_name'], sku['ram_raw'], sku['storage_raw'])
        wanted = (raw['产品名称'], raw['运行内存'], raw['存储'])
        specs = [dict(s) for s in con.execute("SELECT * FROM specifications WHERE (subject_type='sku' AND subject_id=?) OR (subject_type='variant' AND subject_id=?) OR (subject_type='model' AND subject_id=?)", (sku['sku_id'], sku['variant_id'], sku['model_id']))]
        spec_checks = [{'spec_id': s['spec_id'], 'key': s['spec_key'], 'raw': s['raw_value'],
                       'parsed': s['value_num'], 'parse_status': s['parse_status'], 'source_record_id': s['source_record_id'],
                       'raw_matches_selected_config': s['raw_value'] == value_to_text(raw[s['source_col']])}
                       for s in specs if s['source_col'] in raw]
        results.append({'sku_id': sku['sku_id'], 'source_record_id': sku['source_record_id'],
                        'identity_matches': identity == wanted, 'source_row': raw,
                        'spec_checks': spec_checks, 'human_status': 'NOT_CHECKED', 'image_review': 'NOT_CHECKED'})
    con.close()
    wb.close()
    passed = len(results) == 20 and all(r['identity_matches'] and all(s['raw_matches_selected_config'] for s in r['spec_checks']) for r in results)
    report = {'machine_status': 'PASS' if passed else 'FAIL', 'count': len(results),
              'human_status': 'NOT_CHECKED', 'wps_status': 'NOT_CHECKED', 'records': results}
    path = ROOT / 'local_evidence/pilot_review.json'
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"Pilot identity/raw comparisons: {report['machine_status']} ({len(results)} records); human/WPS pending")
    raise SystemExit(0 if passed else 1)
