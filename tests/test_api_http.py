import csv
import importlib.util
from io import BytesIO, StringIO
import json
from pathlib import Path
import socket
import tempfile
import threading
import time
import unittest
import urllib.request
from unittest.mock import patch

from starlette.testclient import TestClient
from tests.api_fixture import make_release, request_payload
from catalog.release import ReleaseStore


class HttpTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec('api.app'),'ASGI app missing')
        from api.app import create_app
        self.temp=tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        root=Path(self.temp.name); self.db=make_release(root)
        static=root/'static'; static.mkdir(); (static/'index.html').write_text('<h1>测试前端</h1>',encoding='utf-8')
        self.app=create_app(self.db,static); self.client=TestClient(self.app)
        self.client.__enter__(); self.addCleanup(self.client.__exit__,None,None,None)
        self.body=request_payload(ReleaseStore(self.db).get())

    def post(self,path='/api/v1/recommend',body=None):
        return self.client.post(path,json=self.body if body is None else body)

    def test_all_read_routes_and_recommend_parity(self):
        self.assertEqual(self.client.get('/healthz').json(),{'status':'alive'})
        self.assertEqual(self.client.get('/readyz').status_code,200)
        meta=self.client.get('/api/v1/meta').json(); self.assertEqual(meta['counts']['skus'],2)
        self.assertEqual(meta['acceptance']['overall_acceptance'],'PENDING')
        listing=self.client.get('/api/v1/catalog').json(); self.assertEqual(listing['page_size'],24)
        self.assertEqual(listing['total'],2)
        self.assertEqual(self.client.get('/api/v1/skus/S1').json()['sku']['sku_id'],'S1')
        compare=self.post('/api/v1/compare',{'sku_ids':['S2','S1']})
        self.assertEqual([r['sku']['sku_id'] for r in compare.json()['items']],['S2','S1'])
        ranked=self.post('/api/v1/recommend?page_size=1').json()
        from api.contracts import parse_recommendation
        result=self.app.state.service.recommend(parse_recommendation(self.body,meta['brands'])).result
        self.assertEqual(ranked['ranking_sha256'],result.ranking_sha256)
        self.assertEqual(ranked['ranked_total'],2); self.assertEqual(len(ranked['ranked']),1)

    def test_request_body_limit(self):
        for count,status in ((32768,400),(32769,413)):
            response=self.client.post('/api/v1/recommend',content=b' '*count,headers={'Content-Type':'application/json'})
            self.assertEqual(response.status_code,status)

    def test_streamed_body_limit_and_forged_length(self):
        response=self.client.post('/api/v1/recommend',content=iter([b' '*16000,b' '*16769]),headers={'Content-Length':'1'})
        self.assertEqual(response.status_code,413)
        response=self.client.post('/api/v1/recommend',content=iter([b' '*20000,b' '*20000]))
        self.assertEqual(response.status_code,413)

    def test_invalid_json_numbers_queries_and_duplicate_keys(self):
        for raw in (b'\xff',b'{',b'{"weights":NaN}',b'{"a":1,"a":2}'):
            self.assertEqual(self.client.post('/api/v1/recommend',content=raw).status_code,400)
        for query in ('page=1&page=2','page=0','page_size=101','db_path=secret'):
            self.assertEqual(self.post('/api/v1/recommend?'+query).status_code,422)
        self.assertEqual(self.post(body={**self.body,'target_screen':True}).status_code,422)
        self.assertEqual(self.post('/api/v1/compare',{'sku_ids':['S1']*5}).status_code,422)
        self.assertEqual(self.client.get('/api/v1/skus/NOPE').status_code,404)

    def test_invalid_release_health_recovers_and_hides_path(self):
        mp=self.db.with_name('release_manifest.json'); original=mp.read_bytes(); mp.write_bytes(b'{}')
        self.assertEqual(self.client.get('/healthz').status_code,200)
        for path in ('/readyz','/api/v1/meta','/api/v1/catalog'):
            r=self.client.get(path); self.assertEqual(r.status_code,503)
            self.assertEqual(r.json()['error']['code'],'release_unavailable'); self.assertNotIn(self.temp.name,r.text)
        mp.write_bytes(original); self.assertEqual(self.client.get('/readyz').status_code,200)

    def test_static_files_and_safe_errors(self):
        self.assertIn('测试前端',self.client.get('/').text)
        for path in ('/api/nope','/data/catalog.sqlite','/.env','/%2e%2e/data/catalog.sqlite'):
            self.assertEqual(self.client.get(path).status_code,404)
        with patch.object(self.app.state.service,'catalog',side_effect=RuntimeError('SECRET /private/path')):
            r=self.client.get('/api/v1/catalog')
        self.assertEqual(r.status_code,500); self.assertNotIn('SECRET',r.text)
        self.assertEqual(r.json()['error']['code'],'internal_error')
        r=self.client.get('/api/v1/meta',headers={'Origin':'https://untrusted.example'})
        self.assertNotIn('access-control-allow-origin',r.headers)

    def test_export_matches_full_result(self):
        p=self.post('/api/v1/recommend?page_size=1').json()
        body={'request':self.body,'fingerprints':{k:p[k] for k in ('decision_sha256','ranking_sha256')}}
        response=self.post('/api/v1/export/csv',body)
        self.assertEqual(response.status_code,200); self.assertTrue(response.content.startswith(b'\xef\xbb\xbf'))
        self.assertIn('attachment;',response.headers['content-disposition'])
        rows=list(csv.DictReader(StringIO(response.content.decode('utf-8-sig')))); self.assertEqual(len(rows),2)
        self.assertEqual({r['sku_id'] for r in rows},{'S1','S2'})
        self.assertEqual(rows[0]['display_name'],'测试机1')
        self.assertEqual(rows[0]['ranking_sha256'],p['ranking_sha256'])
        response=self.post('/api/v1/export/json',body); self.assertEqual(response.status_code,200)
        self.assertIsNone(response.json()['ranking'][0]['dxomark_photo'])
        from openpyxl import load_workbook
        response=self.post('/api/v1/export/xlsx',body); wb=load_workbook(BytesIO(response.content))
        ws=wb['推荐结果']; self.assertEqual(ws['A2'].value,'S1'); self.assertEqual(ws['A2'].number_format,'@')
        col=[c.value for c in ws[1]].index('dxomark_photo')+1; self.assertIsNone(ws.cell(2,col).value); wb.close()
        body['fingerprints']['ranking_sha256']='0'*64
        self.assertEqual(self.post('/api/v1/export/csv',body).status_code,409)

    def test_real_http_roundtrip(self):
        import uvicorn
        sock=socket.socket(); sock.bind(('127.0.0.1',0)); port=sock.getsockname()[1]; sock.close()
        server=uvicorn.Server(uvicorn.Config(self.app,host='127.0.0.1',port=port,log_level='error'))
        thread=threading.Thread(target=server.run,daemon=True); thread.start()
        def stop():server.should_exit=True; thread.join(10)
        self.addCleanup(stop)
        deadline=time.monotonic()+10
        while not server.started and time.monotonic()<deadline:time.sleep(.05)
        self.assertTrue(server.started)
        base=f'http://127.0.0.1:{port}'
        req=urllib.request.Request(base+'/api/v1/recommend',data=json.dumps(self.body).encode(),headers={'Content-Type':'application/json'})
        with urllib.request.urlopen(req,timeout=10) as response:p=json.load(response)
        self.assertEqual(p['ranked_total'],2)
        body={'request':self.body,'fingerprints':{k:p[k] for k in ('decision_sha256','ranking_sha256')}}
        req=urllib.request.Request(base+'/api/v1/export/json',data=json.dumps(body).encode(),headers={'Content-Type':'application/json'})
        with urllib.request.urlopen(req,timeout=10) as response:export=json.load(response)
        self.assertEqual(len(export['ranking']),2); self.assertEqual(export['evidence']['ranking_sha256'],p['ranking_sha256'])

    def test_formula_text_in_http_export_is_display_escaped_only(self):
        from api.app import create_app
        def prepare(con):con.execute("UPDATE phone_models SET display_name='=2+2' WHERE model_id='M1'")
        db=make_release(Path(self.temp.name)/'formula',prepare=prepare)
        with TestClient(create_app(db)) as client:
            request=request_payload(ReleaseStore(db).get())
            p=client.post('/api/v1/recommend',json=request).json()
            body={'request':request,'fingerprints':{k:p[k] for k in ('decision_sha256','ranking_sha256')}}
            exported=client.post('/api/v1/export/csv',json=body)
            rows=list(csv.DictReader(StringIO(exported.content.decode('utf-8-sig'))))
            self.assertEqual(rows[0]['display_name'],"'=2+2")
            self.assertEqual(client.get('/api/v1/skus/S1').json()['model']['display_name'],'=2+2')
