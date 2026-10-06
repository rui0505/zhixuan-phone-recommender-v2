"""Synthetic results: export must preserve values/IDs and reject mixed versions."""
from pathlib import Path
import tempfile
import unittest
from io import BytesIO, StringIO
import csv
import json
from catalog_fixture import create_catalog
from catalog.repository import CatalogRepository
from catalog.evidence import build_evidence
from engine import rank_decision_input


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.db=Path(self.temp.name)/'c.sqlite'
        self.con=create_catalog(self.db)
        self.con.execute("delete from specifications where spec_id='weight_g1'");self.con.commit()
        self.di=CatalogRepository(str(self.db)).build_decision_input(['price_cny','weight_g'],weights={'price_cny':1,'weight_g':0})
        self.result=rank_decision_input(self.di);self.bundle=build_evidence(self.di,self.result,{})
    def tearDown(self):self.con.close();self.temp.cleanup()
    def export(self,fmt):
        from ui.export_results import export_results
        return export_results(self.result,self.bundle,fmt)
    def test_csv_utf8_bom_and_null_preserved(self):
        data=self.export('csv');self.assertTrue(data.startswith(b'\xef\xbb\xbf'))
        rows=list(csv.DictReader(StringIO(data.decode('utf-8-sig'))))
        self.assertEqual(rows[0]['sku_id'],'S1');self.assertEqual(rows[0]['weight_g'],'')
        self.assertEqual(rows[0]['price_cny'],'1000.0')
        self.assertEqual(rows[0]['dataset_version'],'fixture-v1')
    def test_xlsx_id_is_text(self):
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(self.export('xlsx')));sheet=wb['推荐结果']
        self.assertEqual(sheet['A2'].value,'S1');self.assertEqual(sheet['A2'].number_format,'@')
        self.assertEqual(sheet['A2'].data_type,'s')
        col=[c.value for c in sheet[1]].index('weight_g')+1
        self.assertIsNone(sheet.cell(2,col).value);wb.close()
    def test_formula_like_text_export_is_safe(self):
        self.con.execute("update phone_models set display_name='=2+2' where model_id='M1'");self.con.commit()
        self.di=CatalogRepository(str(self.db)).build_decision_input(['price_cny'])
        self.result=rank_decision_input(self.di);self.bundle=build_evidence(self.di,self.result,{})
        rows=list(csv.DictReader(StringIO(self.export('csv').decode('utf-8-sig'))))
        self.assertEqual(rows[0]['display_name'],"'=2+2")
        self.assertEqual(self.con.execute("select display_name from phone_models where model_id='M1'").fetchone()[0],'=2+2')
    def test_json_contains_null_not_nan_and_evidence(self):
        raw=self.export('json').decode('utf-8');data=json.loads(raw)
        self.assertNotIn('NaN',raw);self.assertIsNone(data['ranking'][0]['weight_g'])
        self.assertEqual(data['evidence']['dataset_version'],'fixture-v1')
    def test_other_version_bundle_cannot_export(self):
        self.result.dataset_version='wrong'
        with self.assertRaises(ValueError):self.export('csv')
    def test_xlsx_evidence_is_field_per_row_not_clipped_json(self):
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(self.export('xlsx')))
        self.assertIn('逐项事实',wb.sheetnames)
        sheet=wb['逐项事实']
        self.assertEqual([c.value for c in sheet[1]],['事实 ID','字段','内容'])
        rows=list(sheet.values)[1:]
        self.assertTrue(any(r[1]=='metric_key' and r[2]=='price_cny' for r in rows))
        self.assertTrue(all(len(str(r[2] or ''))<=160 for r in rows))
        wb.close()
