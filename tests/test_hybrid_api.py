"""Consumer contracts for fixed confirmed weights and honest performance limits."""
from copy import deepcopy
import csv
import importlib.util
from io import BytesIO, StringIO
import json
from pathlib import Path
import tempfile
import unittest

import numpy as np
import pandas as pd
from starlette.testclient import TestClient

from api.app import create_app
from api.contracts import parse_recommendation
from catalog.adapter import build_input, decision_input_from_frame
from catalog.binding import verify_result
from catalog.evidence import build_evidence, verify_bundle
from catalog.release import ReleaseStore, file_sha
from engine import rank_decision_input
from tests.api_fixture import make_release, request_payload
from tests.catalog_fixture import insert


class HybridApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.db = make_release(Path(self.temp.name))
        self.release = ReleaseStore(self.db).get()
        self.app = create_app(self.db)
        self.client = TestClient(self.app)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.body = request_payload(self.release)

    def post(self, body=None, path='/api/v1/recommend'):
        return self.client.post(path, json=self.body if body is None else body)

    def hybrid(self, **updates):
        return {**self.body, 'method': 'hybrid_rrf', **updates}

    def test_old_request_and_explicit_default_keep_original_hashes(self):
        # Catch accidental routing of default TOPSIS through the fusion adapter.
        request = parse_recommendation(self.body, ['品牌1', '品牌2'])
        di = build_input(self.release.repo, request.needs, request.weights, 6.5, 'new')
        expected = rank_decision_input(di, target_screen=6.5)
        for body in (self.body, {**self.body, 'method': 'topsis',
                                'expected_algorithm_version': 'topsis-1.0.0'}):
            response = self.post(body)
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertIn('method', payload)
            self.assertEqual(payload['method'], 'topsis')
            self.assertEqual(payload['score_label'], 'TOPSIS 贴近度')
            self.assertEqual(payload['ranking_sha256'], expected.ranking_sha256)
            self.assertEqual(payload['decision_sha256'], expected.decision_sha256)
            self.assertEqual([r['score'] for r in payload['ranked']], expected.ranked['score'].tolist())

    def test_hybrid_is_explicit_versioned_and_bound_to_evidence(self):
        response = self.post(self.hybrid(expected_algorithm_version='hybrid-rrf-1.0.0-experimental'))
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload['method'], 'hybrid_rrf')
        self.assertEqual(payload['score_label'], '融合排序分')
        self.assertEqual(payload['algorithm_version'], 'hybrid-rrf-1.0.0-experimental')
        context = payload['evidence']['decision_context']
        self.assertEqual(context['method'], 'hybrid_rrf')
        self.assertEqual(context['algorithm_parameters'], {'theta': 1, 'alpha': .5, 'k': 10})
        verify_bundle(payload['evidence'])
        self.assertIn('融合排序分', next(iter(payload['explanations'].values())))
        self.assertIn('尚未证明更准确', ' '.join(payload['evidence']['limitations']))
        baseline = self.post().json()
        self.assertEqual(payload['decision_sha256'], baseline['decision_sha256'])
        self.assertNotEqual(payload['ranking_sha256'], baseline['ranking_sha256'])

    def test_invalid_method_and_version_are_rejected(self):
        for value in ('unknown', '', None, True, [], {}):
            with self.subTest(method=value):
                self.assertEqual(self.post({**self.body, 'method': value}).status_code, 422)
        for method in ('topsis', 'hybrid_rrf'):
            response = self.post({**self.body, 'method': method, 'expected_algorithm_version': 'obsolete'})
            self.assertEqual(response.status_code, 409, response.text)
            self.assertEqual(response.json()['error']['code'], 'algorithm_changed')
        for version in (None, True, '', [], 'x' * 257):
            self.assertEqual(self.post({**self.body, 'expected_algorithm_version': version}).status_code, 422)
        self.assertEqual(self.post({**self.hybrid(), 'theta': 2}).status_code, 422)

    def test_changed_method_cannot_export_or_explain_old_fingerprints(self):
        old = self.post().json()
        envelope = {'request': self.hybrid(), 'fingerprints': {k: old[k] for k in
                    ('decision_sha256', 'ranking_sha256')}}
        for fmt in ('csv', 'xlsx', 'json'):
            self.assertEqual(self.post(envelope, '/api/v1/export/' + fmt).status_code, 409)
        response = self.post({**envelope, 'sku_id': old['ranked'][0]['sku_id']}, '/api/v1/explain')
        self.assertEqual(response.status_code, 409)

    def test_hybrid_pagination_exports_and_explanation_are_one_decision(self):
        request = self.hybrid()
        payload = self.post(request).json()
        self.assertIn('ranking_sha256', payload)
        page = self.post(request, '/api/v1/recommend?page=2&page_size=1').json()
        self.assertEqual(page['ranking_sha256'], payload['ranking_sha256'])
        self.assertEqual(page['ranked'][0]['sku_id'], payload['ranked'][1]['sku_id'])
        envelope = {'request': request, 'fingerprints': {k: payload[k] for k in
                    ('decision_sha256', 'ranking_sha256')}}
        exported = self.post(envelope, '/api/v1/export/json')
        self.assertEqual(exported.status_code, 200, exported.text)
        rows = exported.json()['ranking']
        self.assertEqual([r['sku_id'] for r in rows], [r['sku_id'] for r in payload['ranked']])
        self.assertTrue(all(r['method'] == 'hybrid_rrf' and r['score_label'] == '融合排序分' for r in rows))
        csv_response = self.post(envelope, '/api/v1/export/csv')
        csv_rows = list(csv.DictReader(StringIO(csv_response.content.decode('utf-8-sig'))))
        self.assertEqual([float(r['score']) for r in csv_rows], [r['score'] for r in rows])
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(self.post(envelope, '/api/v1/export/xlsx').content))
        self.addCleanup(workbook.close)
        columns = [cell.value for cell in workbook['推荐结果'][1]]
        self.assertEqual(workbook['推荐结果'].cell(2, columns.index('method') + 1).value, 'hybrid_rrf')
        explanation = self.post({**envelope, 'sku_id': rows[0]['sku_id']}, '/api/v1/explain')
        self.assertEqual(explanation.status_code, 200, explanation.text)
        self.assertEqual(explanation.json()['method'], 'hybrid_rrf')
        self.assertEqual(explanation.json()['algorithm_version'], payload['algorithm_version'])
        self.assertIn('融合排序分', explanation.json()['text'])

    def test_battery_zero_is_ignored_but_positive_is_explicitly_unavailable(self):
        for method in ('topsis', 'hybrid_rrf'):
            base = {**self.body, 'method': method}
            expected = self.post(base).json()
            zero = self.post({**base, 'weights': {**self.body['weights'], 'battery_mixed_usage_duration_min': 0}})
            self.assertEqual(zero.status_code, 200, zero.text)
            self.assertEqual(zero.json()['ranking_sha256'], expected['ranking_sha256'])
            positive = self.post({**base, 'weights': {'price_cny': 1, 'battery_mixed_usage_duration_min': 2}})
            self.assertEqual(positive.status_code, 422)
            self.assertEqual(positive.json()['error']['code'], 'unsupported_metric')

    def test_game_remains_strictly_insufficient_in_both_modes(self):
        for method in ('topsis', 'hybrid_rrf'):
            response = self.post({**self.body, 'method': method, 'weights': {'honor_of_kings_avg_fps': 10}})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['status'], 'insufficient_comparable_data')
            self.assertEqual(response.json()['ranked'], [])
            self.assertIsNone(response.json()['evidence'])

    def test_meta_capabilities_respect_filters_and_keep_release_four_fields(self):
        meta = self.client.get('/api/v1/meta').json()
        self.assertIn('algorithms', meta)
        self.assertEqual(set(meta['release']), {'dataset_version', 'feature_rule_version', 'snapshot_sha256', 'manifest_sha256'})
        self.assertEqual({a['method'] for a in meta['algorithms']}, {'topsis', 'hybrid_rrf'})
        capabilities = meta['performance_capabilities']
        self.assertEqual(capabilities['battery_mixed_usage_duration_min']['status'], 'unsupported_metric')
        self.assertFalse(capabilities['honor_of_kings_avg_fps']['available'])
        self.assertEqual(capabilities['honor_of_kings_avg_fps']['eligible_model_count'], 0)
        from api.performance import performance_capabilities
        needs = deepcopy(self.body['needs']); needs['budget']['max'] = 1
        filtered = performance_capabilities(self.release.repo, needs, 'new')
        self.assertEqual(filtered['dxomark_photo']['eligible_model_count'], 0)

    def test_recommendation_capabilities_are_for_current_filters_not_whole_catalog(self):
        # Catch using cached whole-catalog coverage after budget/brand changes.
        from api.performance import performance_capabilities
        def prepare(con):
            for n in (1, 2):
                insert(con, 'benchmark_runs', run_id=f'PHOTO{n}', model_id=f'M{n}',
                       subject_type='model', subject_id=f'M{n}', benchmark='DXOMARK',
                       institution='DXOMARK', protocol='dxomark_v5', applicable_scope='model')
                insert(con, 'metric_observations', observation_id=f'PHOTO{n}', run_id=f'PHOTO{n}',
                       metric_key='dxomark_photo', value_num=100+n, comparison_group='dxomark_v5')
        db = make_release(Path(self.temp.name) / 'photos', prepare=prepare)
        release = ReleaseStore(db).get()
        needs = deepcopy(self.body['needs'])
        self.assertTrue(performance_capabilities(release.repo, needs, 'new')['dxomark_photo']['available'])
        needs['budget']['max'] = 1100
        with TestClient(create_app(db)) as client:
            payload = client.post('/api/v1/recommend', json=request_payload(release, needs=needs)).json()
        self.assertIn('performance_capabilities', payload)
        photo = payload['performance_capabilities']['dxomark_photo']
        self.assertFalse(photo['available'])
        self.assertEqual(photo['eligible_model_count'], 1)
        self.assertEqual(photo['comparison_group'], 'dxomark_v5')

    def test_default_catalog_default_parity_and_hybrid_original_indices(self):
        store = ReleaseStore(Path(__file__).resolve().parents[1] / 'data/catalog.sqlite')
        release = store.get(); before = file_sha(store.db_path)
        from api.service import CatalogService
        service = CatalogService(store)
        body = request_payload(release); brands = service.meta()['brands']
        default = service.recommend(parse_recommendation(body, brands))
        self.assertIn('algorithms', service.meta())
        hybrid = service.recommend(parse_recommendation({**body, 'method': 'hybrid_rrf'}, brands))
        core = rank_decision_input(default.di, target_screen=6.5)
        self.assertEqual(default.result.ranking_sha256, core.ranking_sha256)
        verify_result(hybrid.result)
        verify_bundle(build_evidence(hybrid.di, hybrid.result, hybrid.needs))
        self.assertEqual(set(hybrid.result.ranked['sku_id']), set(core.ranked['sku_id']))
        for row in hybrid.result.ranked.to_dict('records'):
            self.assertEqual(hybrid.di['candidates'].iloc[row['input_index']]['sku_id'], row['sku_id'])
        self.assertEqual(file_sha(store.db_path), before)


class HybridMathTests(unittest.TestCase):
    def mode(self):
        self.assertIsNotNone(importlib.util.find_spec('ranking_modes'), 'production ranking adapter missing')
        import ranking_modes
        return ranking_modes

    def test_crossed_two_metric_classic_todim_and_rrf_midranks(self):
        # Hand derivation: with equal weights gains/losses cancel symmetrically;
        # TOPSIS=.5, TODIM=.5, each midrank=1.5, RRF=1/11.5=2/23.
        mode = self.mode()
        di = decision_input_from_frame(pd.DataFrame([
            {'id': 'Z', '产品名称': 'A', 'price': 100, 'storage': 100},
            {'id': 'A', '产品名称': 'B', 'price': 200, 'storage': 200},
        ]), {'price': 1, 'storage': 1})
        result = mode.rank(di, method='hybrid_rrf')
        np.testing.assert_allclose(result.ranked['score'], [2 / 23, 2 / 23], atol=1e-12)
        np.testing.assert_allclose(result.ranked['todim_score'], [.5, .5], atol=1e-12)
        self.assertEqual(result.ranked['sku_id'].tolist(), ['A', 'Z'])
        self.assertEqual(result.ranked['input_index'].tolist(), [1, 0])
        self.assertIn('并列', result.reason)

    def test_original_indices_survive_backups_and_target_transform(self):
        mode = self.mode()
        di = decision_input_from_frame(pd.DataFrame([
            {'id': 'Z', '产品名称': 'A', 'screen': 6.1},
            {'id': 'MISSING', '产品名称': 'B', 'screen': None},
            {'id': 'A', '产品名称': 'C', 'screen': 6.9},
        ]), {'screen': 1})
        result = mode.rank(di, target_screen=6.1, method='hybrid_rrf')
        self.assertEqual(result.ranked['sku_id'].tolist(), ['Z', 'A'])
        self.assertEqual(result.ranked['input_index'].tolist(), [0, 2])
        self.assertEqual(result.backups['sku_id'].tolist(), ['MISSING'])
        self.assertNotIn('score', result.backups)
        verify_result(result)
        with self.assertRaises(ValueError): mode.rank(di, method='hybrid_rrf')

    def test_classic_inverse_weight_loss_counterexample_is_retained(self):
        mode = self.mode()
        di = decision_input_from_frame(pd.DataFrame([
            {'id': 'A', '产品名称': 'A', 'storage': 100, 'refresh': 60},
            {'id': 'B', '产品名称': 'B', 'storage': 200, 'refresh': 30},
        ]), {'storage': 9, 'refresh': 1})
        result = mode.rank(di, method='hybrid_rrf')
        scores = result.ranked.set_index('sku_id')
        self.assertGreater(scores.loc['B', 'topsis_score'], scores.loc['A', 'topsis_score'])
        self.assertGreater(scores.loc['A', 'todim_score'], scores.loc['B', 'todim_score'])
        np.testing.assert_allclose(scores['score'], [23 / 264, 23 / 264], atol=1e-12)

